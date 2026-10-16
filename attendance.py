"""
attendance.py — attendance data layer.

Two stores are maintained in parallel:
  1. attendance.csv  — raw append-only log (all completed entries)
  2. master_attendance.xlsx — single print-ready master sheet; updated live
     on every /present (row added, Status=Present) and /leave (row filled).

CSV columns:
    User_ID, Username, Full_Name,
    Join_Date, Join_Time, Leave_Date, Leave_Time,
    Duration_Minutes, Session, Chat_ID

Master Excel columns (print sheet):
    No | User ID | Full Name | Date | Present Time | Leave Time | Duration | Status
"""

import os
from datetime import datetime, timezone
from typing import Optional

import pandas as pd
from openpyxl import Workbook, load_workbook
from openpyxl.styles import (
    Alignment, Border, Font, PatternFill, Side
)
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.page import PageMargins

from config import CSV_FILE, MASTER_XLSX, TIMEZONE

# ── CSV columns ───────────────────────────────────────────────────────────────
COLUMNS = [
    "User_ID", "Username", "Full_Name",
    "Join_Date", "Join_Time",
    "Leave_Date", "Leave_Time",
    "Duration_Minutes", "Session", "Chat_ID",
]

# ── Master Excel columns (display order) ─────────────────────────────────────
MASTER_COLS = [
    "No", "User_ID", "Full_Name",
    "Date", "Present_Time", "Leave_Time",
    "Duration", "Status",
]

# In-memory active store
# { chat_id: { user_id: { username, name, join_time, session } } }
_active: dict[int, dict[int, dict]] = {}

# ── styling constants ─────────────────────────────────────────────────────────
_TITLE_FONT   = Font(name="Calibri", bold=True, size=13)
_HEADER_FONT  = Font(name="Calibri", bold=True, size=11, color="FFFFFF")
_DATA_FONT    = Font(name="Calibri", size=10)
_HEADER_FILL  = PatternFill("solid", fgColor="4472C4")   # blue
_PRESENT_FILL = PatternFill("solid", fgColor="E2EFDA")   # light green
_LEFT_FILL    = PatternFill("solid", fgColor="FCE4D6")   # light orange
_ALT_FILL     = PatternFill("solid", fgColor="F2F2F2")   # light grey
_CENTER       = Alignment(horizontal="center", vertical="center", wrap_text=False)
_LEFT_ALIGN   = Alignment(horizontal="left",   vertical="center", wrap_text=False)

_THIN  = Side(style="thin",  color="AAAAAA")
_MED   = Side(style="medium", color="444444")
_THIN_BORDER   = Border(left=_THIN, right=_THIN, top=_THIN, bottom=_THIN)
_HEADER_BORDER = Border(left=_MED,  right=_MED,  top=_MED,  bottom=_MED)

# Column definitions: (header label, width, alignment)
_MASTER_COLS = [
    ("No",           6,  _CENTER),
    ("User ID",      14, _CENTER),
    ("Full Name",    28, _LEFT_ALIGN),
    ("Date",         13, _CENTER),
    ("Present Time", 14, _CENTER),
    ("Leave Time",   12, _CENTER),
    ("Duration",     13, _CENTER),
    ("Status",       11, _CENTER),
]


# ── helpers ───────────────────────────────────────────────────────────────────

def _now() -> datetime:
    return datetime.now(tz=timezone.utc).astimezone(TIMEZONE)


def _dur_str(minutes: float) -> str:
    total_secs = int(minutes * 60)
    h, rem = divmod(total_secs, 3600)
    m, s   = divmod(rem, 60)
    if h:
        return f"{h}h {m}m {s}s"
    return f"{m}m {s}s"


# ── CSV helpers ───────────────────────────────────────────────────────────────

def init_csv() -> None:
    if not os.path.exists(CSV_FILE):
        pd.DataFrame(columns=COLUMNS).to_csv(CSV_FILE, index=False)


def _read_csv() -> pd.DataFrame:
    init_csv()
    return pd.read_csv(CSV_FILE, dtype={"User_ID": "Int64", "Chat_ID": "Int64"})


def _append_csv(row: dict) -> None:
    init_csv()
    join_dt:  datetime = row["join_time"]
    leave_dt: datetime = row["leave_time"]
    entry = {
        "User_ID":          row["user_id"],
        "Username":         row["username"] or "",
        "Full_Name":        row["name"],
        "Join_Date":        join_dt.strftime("%Y-%m-%d"),
        "Join_Time":        join_dt.strftime("%H:%M:%S"),
        "Leave_Date":       leave_dt.strftime("%Y-%m-%d"),
        "Leave_Time":       leave_dt.strftime("%H:%M:%S"),
        "Duration_Minutes": round(row["duration_minutes"], 2),
        "Session":          row["session"],
        "Chat_ID":          row["chat_id"],
    }
    df = pd.read_csv(CSV_FILE)
    df = pd.concat([df, pd.DataFrame([entry])], ignore_index=True)
    df.to_csv(CSV_FILE, index=False)


# ── Master Excel helpers ──────────────────────────────────────────────────────

def _set_print_props(ws) -> None:
    """Configure worksheet for clean landscape printing on A4."""
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToPage   = True
    ws.page_setup.fitToWidth  = 1
    ws.page_setup.fitToHeight = 0          # unlimited height
    ws.page_setup.paperSize   = ws.PAPERSIZE_A4
    ws.print_title_rows       = "1:2"      # repeat title + header on every page
    ws.page_margins = PageMargins(
        left=0.5, right=0.5, top=0.75, bottom=0.75,
        header=0.3, footer=0.3,
    )
    ws.sheet_view.showGridLines = True


def _write_title_row(ws) -> None:
    """
    Row 1 — plain unmerged title cell in A1 only.
    No merge_cells; the text sits in A1 and the empty B1-H1 cells are just blank.
    """
    c = ws.cell(1, 1, "Student / Member Attendance Summary Sheet")
    c.font      = _TITLE_FONT
    c.alignment = _LEFT_ALIGN
    ws.row_dimensions[1].height = 24


def _write_header_row(ws) -> None:
    """Row 2 — bold blue header, one cell per column, no merging."""
    for col, (label, width, align) in enumerate(_MASTER_COLS, 1):
        c = ws.cell(2, col, label)
        c.font      = _HEADER_FONT
        c.fill      = _HEADER_FILL
        c.alignment = _CENTER        # headers always centred
        c.border    = _HEADER_BORDER
        ws.column_dimensions[get_column_letter(col)].width = width
    ws.row_dimensions[2].height = 18


def _write_data_row(ws, excel_row: int, seq_no: int, user_id, name: str,
                    date_str: str, present_time: str, leave_time: str,
                    duration: str, status: str) -> None:
    """Write one unmerged data row at excel_row."""
    row_fill = _PRESENT_FILL if status == "Present" else (
        _LEFT_FILL if seq_no % 2 != 0 else _ALT_FILL
    )
    vals = [seq_no, user_id, name, date_str, present_time, leave_time, duration, status]
    for col, (val, (_, _, align)) in enumerate(zip(vals, _MASTER_COLS), 1):
        c = ws.cell(excel_row, col, val)
        c.font      = _DATA_FONT
        c.alignment = align
        c.border    = _THIN_BORDER
        c.fill      = row_fill
    ws.row_dimensions[excel_row].height = 16


def _rebuild_master() -> None:
    """
    Rebuild master_attendance.xlsx from scratch using the CSV as source of truth.
    Called after a session delete. No merged cells anywhere.
    """
    df = _read_csv()
    wb = Workbook()
    ws = wb.active
    ws.title = "Attendance"

    _write_title_row(ws)
    _write_header_row(ws)
    _set_print_props(ws)

    for i, (_, row) in enumerate(df.iterrows(), 1):
        dur_raw = row.get("Duration_Minutes", 0)
        _write_data_row(
            ws, excel_row=i + 2, seq_no=i,
            user_id=int(row["User_ID"]) if pd.notna(row.get("User_ID")) else "",
            name=str(row.get("Full_Name", "")),
            date_str=str(row.get("Join_Date", "")),
            present_time=str(row.get("Join_Time", "")),
            leave_time=str(row.get("Leave_Time", "")),
            duration=_dur_str(float(dur_raw)) if pd.notna(dur_raw) and dur_raw else "—",
            status="Left",
        )

    wb.save(MASTER_XLSX)


def _init_master_wb() -> tuple:
    """
    Load existing master workbook, or create a fresh one with title+header rows.
    Returns (wb, ws, data_row_count).
    No merge_cells used anywhere.
    """
    if os.path.exists(MASTER_XLSX):
        wb = load_workbook(MASTER_XLSX)
        ws = wb.active
        data_rows = max(0, ws.max_row - 2)   # subtract title row + header row
        return wb, ws, data_rows

    # ── Brand-new file ────────────────────────────────────────────────────────
    wb = Workbook()
    ws = wb.active
    ws.title = "Attendance"
    _write_title_row(ws)
    _write_header_row(ws)
    _set_print_props(ws)
    return wb, ws, 0


def _upsert_master(user_id: int, name: str, join_dt: datetime,
                   leave_dt: Optional[datetime], duration_minutes: Optional[float],
                   session: str, status: str) -> None:
    """
    Incrementally update master_attendance.xlsx.
    /present → append a new row with Status=Present, Leave Time and Duration blank.
    /leave   → find the matching Present row and fill Leave Time, Duration, Status.
    Zero merged cells in the entire file.
    """
    wb, ws, data_rows = _init_master_wb()

    if status == "Present":
        # ── Append a new Present row ──────────────────────────────────────────
        seq_no    = data_rows + 1
        excel_row = data_rows + 3   # row 1=title, row 2=header, row 3+=data
        _write_data_row(
            ws, excel_row=excel_row, seq_no=seq_no,
            user_id=user_id, name=name,
            date_str=join_dt.strftime("%Y-%m-%d"),
            present_time=join_dt.strftime("%H:%M:%S"),
            leave_time="—",
            duration="—",
            status="Present",
        )

    else:
        # ── Find matching Present row and fill it in ──────────────────────────
        target_row = None
        join_date  = join_dt.strftime("%Y-%m-%d")
        join_time  = join_dt.strftime("%H:%M:%S")

        for r in range(3, ws.max_row + 1):
            cell_uid    = ws.cell(r, 2).value   # col B = User ID
            cell_date   = ws.cell(r, 4).value   # col D = Date
            cell_ptime  = ws.cell(r, 5).value   # col E = Present Time
            cell_status = ws.cell(r, 8).value   # col H = Status
            if (str(cell_uid) == str(user_id)
                    and str(cell_date) == join_date
                    and str(cell_ptime) == join_time
                    and cell_status == "Present"):
                target_row = r
                break

        dur_disp  = _dur_str(duration_minutes) if duration_minutes is not None else "—"
        leave_str = leave_dt.strftime("%H:%M:%S") if leave_dt else "—"

        if target_row:
            ws.cell(target_row, 6).value = leave_str   # Leave Time
            ws.cell(target_row, 7).value = dur_disp    # Duration
            ws.cell(target_row, 8).value = "Left"      # Status
            # Re-apply fill and border to the whole row
            seq_no = ws.cell(target_row, 1).value or 0
            row_fill = _LEFT_FILL if int(seq_no) % 2 != 0 else _ALT_FILL
            for col in range(1, 9):
                c = ws.cell(target_row, col)
                c.fill   = row_fill
                c.border = _THIN_BORDER
        else:
            # Fallback: Present row not found — append a complete Left row
            seq_no    = data_rows + 1
            excel_row = ws.max_row + 1
            _write_data_row(
                ws, excel_row=excel_row, seq_no=seq_no,
                user_id=user_id, name=name,
                date_str=join_date,
                present_time=join_time,
                leave_time=leave_str,
                duration=dur_disp,
                status="Left",
            )

    wb.save(MASTER_XLSX)


# ── public API ────────────────────────────────────────────────────────────────

def record_join(
    chat_id: int,
    session_label: str,
    user_id: int,
    name: str,
    username: Optional[str],
) -> datetime:
    """Store join in memory and add a 'Present' row to the master Excel."""
    join_time = _now()
    _active.setdefault(chat_id, {})[user_id] = {
        "username": username,
        "name": name,
        "join_time": join_time,
        "session": session_label,
    }
    # Write Present row immediately to master Excel
    _upsert_master(
        user_id=user_id, name=name,
        join_dt=join_time, leave_dt=None,
        duration_minutes=None,
        session=session_label, status="Present",
    )
    return join_time


def record_leave(
    chat_id: int,
    session_label: str,
    user_id: int,
) -> Optional[datetime]:
    """Pop from memory, append to CSV, and update the master Excel row to 'Left'."""
    entry = _active.get(chat_id, {}).pop(user_id, None)
    if entry is None:
        return None

    leave_time      = _now()
    join_time       = entry["join_time"]
    stored_session  = entry.get("session") or session_label
    duration_mins   = (leave_time - join_time).total_seconds() / 60

    # 1. Append to CSV log
    _append_csv({
        "user_id":          user_id,
        "username":         entry["username"],
        "name":             entry["name"],
        "join_time":        join_time,
        "leave_time":       leave_time,
        "duration_minutes": duration_mins,
        "session":          stored_session,
        "chat_id":          chat_id,
    })

    # 2. Update the master Excel row
    _upsert_master(
        user_id=user_id, name=entry["name"],
        join_dt=join_time, leave_dt=leave_time,
        duration_minutes=duration_mins,
        session=stored_session, status="Left",
    )
    return leave_time


def close_all_active(chat_id: int, session_label: str) -> int:
    user_ids = list(_active.get(chat_id, {}).keys())
    for uid in user_ids:
        record_leave(chat_id, session_label, uid)
    return len(user_ids)


def get_session_df(chat_id: int, session_label: str) -> pd.DataFrame:
    df = _read_csv()
    if df.empty:
        return pd.DataFrame(columns=COLUMNS)
    mask = (df["Chat_ID"] == chat_id) & (df["Session"] == session_label)
    return df[mask].copy().reset_index(drop=True)


def list_sessions(chat_id: int) -> list[str]:
    df = _read_csv()
    if df.empty:
        return []
    return df[df["Chat_ID"] == chat_id]["Session"].dropna().unique().tolist()


def delete_session(chat_id: int, session_label: str) -> bool:
    df = _read_csv()
    mask = (df["Chat_ID"] == chat_id) & (df["Session"] == session_label)
    if not mask.any():
        return False
    df = df[~mask].reset_index(drop=True)
    df.to_csv(CSV_FILE, index=False)
    _rebuild_master()   # rebuild master from updated CSV
    return True


def get_active_participants(chat_id: int) -> list[dict]:
    return [
        {"user_id": uid, **info}
        for uid, info in _active.get(chat_id, {}).items()
    ]


def is_active(chat_id: int, user_id: int) -> bool:
    return user_id in _active.get(chat_id, {})
