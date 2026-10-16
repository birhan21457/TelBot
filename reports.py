"""
reports.py — build text and Excel reports from a pandas DataFrame.

DataFrame columns expected (from attendance.get_session_df):
    User_ID, Username, Full_Name, Join_Time, Leave_Time,
    Duration_Minutes, Session, Chat_ID
"""

import html
import os

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from config import EXPORT_DIR


# ── text report ───────────────────────────────────────────────────────────────

def build_text_report(session_label: str, df: pd.DataFrame) -> str:
    """
    Return an HTML-formatted string suitable for a Telegram message.
    Groups rows by participant and shows each join/leave entry.
    """
    lines = [
        f"📋 <b>Attendance Report</b>",
        f"🏷 Session: <code>{html.escape(session_label)}</code>",
        f"👥 Unique participants: <b>{df['User_ID'].nunique()}</b>",
        f"📝 Total entries: <b>{len(df)}</b>",
        "",
    ]

    # Group by participant
    for user_id, group in df.groupby("User_ID"):
        row0 = group.iloc[0]
        name  = html.escape(str(row0["Full_Name"]))
        uname = str(row0["Username"]) if pd.notna(row0["Username"]) and row0["Username"] else ""
        upart = f" {html.escape(uname)}" if uname else ""

        total_mins = group["Duration_Minutes"].sum()
        h, rem = divmod(int(total_mins * 60), 3600)
        m, s   = divmod(rem, 60)
        total_str = f"{h}h {m}m {s}s" if h else f"{m}m {s}s"

        lines.append(f"👤 <b>{name}</b>{upart}")
        lines.append(f"   ⏱ Total time: <b>{total_str}</b>  ({len(group)} visit{'s' if len(group) > 1 else ''})")

        for i, (_, entry) in enumerate(group.iterrows(), 1):
            join_t  = str(entry["Join_Time"])[11:19]   # HH:MM:SS slice
            leave_t = str(entry["Leave_Time"])[11:19]
            dur_m   = float(entry["Duration_Minutes"])
            dm, ds  = divmod(int(dur_m * 60), 60)
            dur_str = f"{dm}m {ds}s"
            lines.append(f"   {i}. {join_t} → {leave_t}  ({dur_str})")

        lines.append("")

    return "\n".join(lines).strip()


# ── Excel report ──────────────────────────────────────────────────────────────

_HDR_FILL  = PatternFill("solid", fgColor="1F4E79")
_HDR_FONT  = Font(color="FFFFFF", bold=True, size=11)
_ALT_FILL  = PatternFill("solid", fgColor="D6E4F0")
_BOLD      = Font(bold=True)
_CENTER    = Alignment(horizontal="center", vertical="center")
_LEFT      = Alignment(horizontal="left",  vertical="center")


def _set_widths(ws, widths: list[int]) -> None:
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w


def build_excel_report(chat_id: int, session_label: str, df: pd.DataFrame) -> str:
    """
    Build a styled Excel workbook with two sheets:
      • Summary  — one row per participant (total duration)
      • Details  — one row per CSV entry (raw Join/Leave/Duration)

    Returns the saved file path.
    """
    wb = Workbook()

    # ── Sheet 1: Summary ──────────────────────────────────────────────────────
    ws = wb.active
    ws.title = "Summary"
    ws.row_dimensions[1].height = 22

    title = ws.cell(1, 1, f"Attendance — {session_label}")
    title.font = Font(bold=True, size=13)
    ws.merge_cells("A1:F1")

    sum_headers = ["#", "Full Name", "Username", "User ID", "Visits", "Total Time (min)"]
    for col, h in enumerate(sum_headers, 1):
        c = ws.cell(2, col, h)
        c.font, c.fill, c.alignment = _HDR_FONT, _HDR_FILL, _CENTER

    summary = (
        df.groupby(["User_ID", "Full_Name", "Username"], as_index=False)
          .agg(Visits=("Duration_Minutes", "count"),
               Total_Minutes=("Duration_Minutes", "sum"))
          .sort_values("Full_Name")
    )

    for i, (_, row) in enumerate(summary.iterrows(), 1):
        fill = _ALT_FILL if i % 2 == 0 else None
        vals = [
            i,
            row["Full_Name"],
            row["Username"] if pd.notna(row["Username"]) and row["Username"] else "—",
            int(row["User_ID"]),
            int(row["Visits"]),
            round(float(row["Total_Minutes"]), 2),
        ]
        for col, v in enumerate(vals, 1):
            c = ws.cell(i + 2, col, v)
            if fill:
                c.fill = fill
            c.alignment = _LEFT

    _set_widths(ws, [5, 28, 20, 14, 8, 18])

    # ── Sheet 2: Details ──────────────────────────────────────────────────────
    wd = wb.create_sheet("Details")
    wd.row_dimensions[1].height = 22

    title2 = wd.cell(1, 1, f"Detailed Log — {session_label}")
    title2.font = Font(bold=True, size=13)
    wd.merge_cells("A1:G1")

    det_headers = ["#", "Full Name", "Username", "Join Time", "Leave Time", "Duration (min)", "User ID"]
    for col, h in enumerate(det_headers, 1):
        c = wd.cell(2, col, h)
        c.font, c.fill, c.alignment = _HDR_FONT, _HDR_FILL, _CENTER

    detail_df = df.sort_values(["Full_Name", "Join_Time"]).reset_index(drop=True)
    for i, (_, row) in enumerate(detail_df.iterrows(), 1):
        fill = _ALT_FILL if i % 2 == 0 else None
        vals = [
            i,
            row["Full_Name"],
            row["Username"] if pd.notna(row["Username"]) and row["Username"] else "—",
            str(row["Join_Time"]),
            str(row["Leave_Time"]),
            round(float(row["Duration_Minutes"]), 2),
            int(row["User_ID"]),
        ]
        for col, v in enumerate(vals, 1):
            c = wd.cell(i + 2, col, v)
            if fill:
                c.fill = fill
            c.alignment = _LEFT

    _set_widths(wd, [5, 28, 20, 20, 20, 16, 14])

    # ── Save ──────────────────────────────────────────────────────────────────
    safe = "".join(c if c.isalnum() or c in " _-" else "_" for c in session_label)
    filename = f"attendance_{chat_id}_{safe[:40]}.xlsx"
    filepath = os.path.join(EXPORT_DIR, filename)
    wb.save(filepath)
    return filepath
